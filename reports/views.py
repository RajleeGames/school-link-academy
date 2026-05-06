from decimal import Decimal

from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import render

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from accounts.decorators import permission_required

from academics.models import ClassLevel
from attendance.models import StudentAttendance, StaffAttendance
from expenses.models import Expense
from fees.models import FeePayment, StudentInvoice
from students.models import Student


# =====================================================
# EXCEL HELPERS
# =====================================================

def style_excel_sheet(ws, title):
    ws.freeze_panes = "A4"

    ws["A1"] = "School Link"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="417690")
    ws["A1"].alignment = Alignment(horizontal="left")

    ws["A2"] = title
    ws["A2"].font = Font(size=12, bold=True, color="2B5668")

    thin = Side(border_style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill = PatternFill("solid", fgColor="417690")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")

    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column

        for cell in column_cells:
            try:
                value_length = len(str(cell.value))
                if value_length > max_length:
                    max_length = value_length
            except Exception:
                pass

        ws.column_dimensions[get_column_letter(column)].width = min(max_length + 4, 35)


def excel_response(workbook, filename):
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def money(value):
    if value is None:
        return 0
    return float(value)


# =====================================================
# REPORT FILTER HELPERS
# =====================================================

def get_filtered_students(request):
    query = request.GET.get("q", "").strip()
    class_filter = request.GET.get("class", "").strip()
    status_filter = request.GET.get("status", "").strip()

    students = Student.objects.select_related(
        "class_level",
        "stream",
        "parent_guardian"
    ).all()

    if query:
        students = students.filter(
            Q(admission_number__icontains=query) |
            Q(first_name__icontains=query) |
            Q(middle_name__icontains=query) |
            Q(last_name__icontains=query) |
            Q(parent_guardian__full_name__icontains=query) |
            Q(parent_guardian__phone__icontains=query)
        )

    if class_filter:
        students = students.filter(class_level_id=class_filter)

    if status_filter:
        students = students.filter(status=status_filter)

    return students, query, class_filter, status_filter


def get_filtered_invoices(request):
    query = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "").strip()
    class_filter = request.GET.get("class", "").strip()

    invoices = StudentInvoice.objects.select_related(
        "student",
        "student__class_level",
        "student__stream",
        "academic_year",
        "term",
    ).all()

    if query:
        invoices = invoices.filter(
            Q(invoice_number__icontains=query) |
            Q(student__admission_number__icontains=query) |
            Q(student__first_name__icontains=query) |
            Q(student__middle_name__icontains=query) |
            Q(student__last_name__icontains=query)
        )

    if status_filter:
        invoices = invoices.filter(status=status_filter)

    if class_filter:
        invoices = invoices.filter(student__class_level_id=class_filter)

    return invoices, query, status_filter, class_filter


def get_filtered_expenses(request):
    query = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()

    expenses = Expense.objects.select_related("category").all()

    if query:
        expenses = expenses.filter(
            Q(expense_number__icontains=query) |
            Q(title__icontains=query) |
            Q(paid_to__icontains=query) |
            Q(reference__icontains=query) |
            Q(category__name__icontains=query)
        )

    if status_filter:
        expenses = expenses.filter(status=status_filter)

    if date_from:
        expenses = expenses.filter(expense_date__gte=date_from)

    if date_to:
        expenses = expenses.filter(expense_date__lte=date_to)

    return expenses, query, status_filter, date_from, date_to


# =====================================================
# NORMAL REPORT PAGES
# =====================================================

@permission_required("reports.view")
def reports_home(request):
    total_students = Student.objects.count()

    total_fees_paid = FeePayment.objects.filter(status="confirmed").aggregate(
        total=Sum("amount")
    )["total"] or Decimal("0.00")

    total_expenses = Expense.objects.filter(status="paid").aggregate(
        total=Sum("amount")
    )["total"] or Decimal("0.00")

    outstanding_fees = Decimal("0.00")
    unpaid_invoices = StudentInvoice.objects.exclude(status__in=["paid", "cancelled"])

    for invoice in unpaid_invoices:
        outstanding_fees += invoice.balance

    context = {
        "total_students": total_students,
        "total_fees_paid": total_fees_paid,
        "total_expenses": total_expenses,
        "outstanding_fees": outstanding_fees,
        "net_cash": total_fees_paid - total_expenses,
    }

    return render(request, "reports/home.html", context)


@permission_required("reports.students")
def student_report(request):
    students, query, class_filter, status_filter = get_filtered_students(request)

    context = {
        "students": students,
        "classes": ClassLevel.objects.filter(is_active=True),
        "query": query,
        "class_filter": class_filter,
        "status_filter": status_filter,
    }

    return render(request, "reports/student_report.html", context)


@permission_required("reports.finance")
def fees_report(request):
    invoices, query, status_filter, class_filter = get_filtered_invoices(request)

    total_invoice_amount = Decimal("0.00")
    total_paid_amount = Decimal("0.00")
    total_balance_amount = Decimal("0.00")

    for invoice in invoices:
        total_invoice_amount += invoice.payable_amount
        total_paid_amount += invoice.paid_amount
        total_balance_amount += invoice.balance

    context = {
        "invoices": invoices,
        "classes": ClassLevel.objects.filter(is_active=True),
        "status_choices": StudentInvoice.STATUS_CHOICES,
        "query": query,
        "status_filter": status_filter,
        "class_filter": class_filter,
        "total_invoice_amount": total_invoice_amount,
        "total_paid_amount": total_paid_amount,
        "total_balance_amount": total_balance_amount,
    }

    return render(request, "reports/fees_report.html", context)


@permission_required("reports.finance")
def expenses_report(request):
    expenses, query, status_filter, date_from, date_to = get_filtered_expenses(request)

    total_expenses = expenses.filter(status="paid").aggregate(
        total=Sum("amount")
    )["total"] or Decimal("0.00")

    pending_expenses = expenses.filter(status="pending").aggregate(
        total=Sum("amount")
    )["total"] or Decimal("0.00")

    context = {
        "expenses": expenses,
        "status_choices": Expense.STATUS_CHOICES,
        "query": query,
        "status_filter": status_filter,
        "date_from": date_from,
        "date_to": date_to,
        "total_expenses": total_expenses,
        "pending_expenses": pending_expenses,
    }

    return render(request, "reports/expenses_report.html", context)


@permission_required("attendance.reports")
def attendance_report(request):
    report_type = request.GET.get("type", "student").strip()
    date_filter = request.GET.get("date", "").strip()
    status_filter = request.GET.get("status", "").strip()

    student_records = StudentAttendance.objects.none()
    staff_records = StaffAttendance.objects.none()

    if report_type == "staff":
        staff_records = StaffAttendance.objects.select_related("staff").all()

        if date_filter:
            staff_records = staff_records.filter(date=date_filter)

        if status_filter:
            staff_records = staff_records.filter(status=status_filter)
    else:
        student_records = StudentAttendance.objects.select_related(
            "student",
            "student__class_level",
            "student__stream",
        ).all()

        if date_filter:
            student_records = student_records.filter(date=date_filter)

        if status_filter:
            student_records = student_records.filter(status=status_filter)

    context = {
        "report_type": report_type,
        "date_filter": date_filter,
        "status_filter": status_filter,
        "student_records": student_records,
        "staff_records": staff_records,
        "student_status_choices": StudentAttendance.STATUS_CHOICES,
        "staff_status_choices": StaffAttendance.STATUS_CHOICES,
    }

    return render(request, "reports/attendance_report.html", context)


# =====================================================
# EXCEL EXPORTS
# =====================================================

@permission_required("reports.students")
def export_student_report_excel(request):
    students, query, class_filter, status_filter = get_filtered_students(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [
        "#",
        "Admission No",
        "Student Name",
        "Gender",
        "Class",
        "Stream",
        "Parent / Guardian",
        "Phone",
        "Status",
    ]

    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(headers)

    for index, student in enumerate(students, start=1):
        ws.append([
            index,
            student.admission_number,
            student.full_name,
            student.get_gender_display(),
            student.class_level.name,
            student.stream.name if student.stream else "",
            student.parent_guardian.full_name,
            student.parent_guardian.phone,
            student.get_status_display(),
        ])

    style_excel_sheet(ws, "Student Report")

    return excel_response(wb, "student_report.xlsx")


@permission_required("reports.finance")
def export_fees_report_excel(request):
    invoices, query, status_filter, class_filter = get_filtered_invoices(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Fees Report"

    headers = [
        "#",
        "Invoice",
        "Student",
        "Admission No",
        "Class",
        "Stream",
        "Academic Year",
        "Term",
        "Payable",
        "Paid",
        "Balance",
        "Status",
    ]

    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(headers)

    total_payable = Decimal("0.00")
    total_paid = Decimal("0.00")
    total_balance = Decimal("0.00")

    for index, invoice in enumerate(invoices, start=1):
        total_payable += invoice.payable_amount
        total_paid += invoice.paid_amount
        total_balance += invoice.balance

        ws.append([
            index,
            invoice.invoice_number,
            invoice.student.full_name,
            invoice.student.admission_number,
            invoice.student.class_level.name,
            invoice.student.stream.name if invoice.student.stream else "",
            invoice.academic_year.name,
            invoice.term.name if invoice.term else "",
            money(invoice.payable_amount),
            money(invoice.paid_amount),
            money(invoice.balance),
            invoice.get_status_display(),
        ])

    ws.append([])
    ws.append([
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "TOTAL",
        money(total_payable),
        money(total_paid),
        money(total_balance),
        "",
    ])

    style_excel_sheet(ws, "Fees Report")

    return excel_response(wb, "fees_report.xlsx")


@permission_required("reports.finance")
def export_expenses_report_excel(request):
    expenses, query, status_filter, date_from, date_to = get_filtered_expenses(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"

    headers = [
        "#",
        "Expense No",
        "Title",
        "Category",
        "Date",
        "Paid To",
        "Method",
        "Reference",
        "Status",
        "Amount",
    ]

    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(headers)

    total_paid = Decimal("0.00")
    total_pending = Decimal("0.00")

    for index, expense in enumerate(expenses, start=1):
        if expense.status == "paid":
            total_paid += expense.amount

        if expense.status == "pending":
            total_pending += expense.amount

        ws.append([
            index,
            expense.expense_number,
            expense.title,
            expense.category.name,
            expense.expense_date,
            expense.paid_to,
            expense.get_payment_method_display(),
            expense.reference,
            expense.get_status_display(),
            money(expense.amount),
        ])

    ws.append([])
    ws.append(["", "", "", "", "", "", "", "TOTAL PAID", "", money(total_paid)])
    ws.append(["", "", "", "", "", "", "", "TOTAL PENDING", "", money(total_pending)])

    style_excel_sheet(ws, "Expenses Report")

    return excel_response(wb, "expenses_report.xlsx")


@permission_required("attendance.reports")
def export_attendance_report_excel(request):
    report_type = request.GET.get("type", "student").strip()
    date_filter = request.GET.get("date", "").strip()
    status_filter = request.GET.get("status", "").strip()

    wb = Workbook()
    ws = wb.active

    if report_type == "staff":
        ws.title = "Staff Attendance"

        headers = [
            "#",
            "Date",
            "Staff",
            "Role",
            "Status",
            "Check In",
            "Check Out",
            "Remarks",
        ]

        staff_records = StaffAttendance.objects.select_related("staff").all()

        if date_filter:
            staff_records = staff_records.filter(date=date_filter)

        if status_filter:
            staff_records = staff_records.filter(status=status_filter)

        ws.append([])
        ws.append([])
        ws.append([])
        ws.append(headers)

        for index, record in enumerate(staff_records, start=1):
            ws.append([
                index,
                record.date,
                record.staff.full_name,
                record.staff.get_role_display(),
                record.get_status_display(),
                record.check_in_time.strftime("%H:%M") if record.check_in_time else "",
                record.check_out_time.strftime("%H:%M") if record.check_out_time else "",
                record.remarks,
            ])

        style_excel_sheet(ws, "Staff Attendance Report")

        return excel_response(wb, "staff_attendance_report.xlsx")

    ws.title = "Student Attendance"

    headers = [
        "#",
        "Date",
        "Student",
        "Admission No",
        "Class",
        "Stream",
        "Status",
        "Arrival",
        "Remarks",
    ]

    student_records = StudentAttendance.objects.select_related(
        "student",
        "student__class_level",
        "student__stream",
    ).all()

    if date_filter:
        student_records = student_records.filter(date=date_filter)

    if status_filter:
        student_records = student_records.filter(status=status_filter)

    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(headers)

    for index, record in enumerate(student_records, start=1):
        ws.append([
            index,
            record.date,
            record.student.full_name,
            record.student.admission_number,
            record.student.class_level.name,
            record.student.stream.name if record.student.stream else "",
            record.get_status_display(),
            record.arrival_time.strftime("%H:%M") if record.arrival_time else "",
            record.remarks,
        ])

    style_excel_sheet(ws, "Student Attendance Report")

    return excel_response(wb, "student_attendance_report.xlsx")