from collections import defaultdict
from decimal import Decimal

from django.contrib import messages
from django.db.models import Avg, Q, Sum
from django.db.models.deletion import ProtectedError
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from accounts.decorators import permission_required
from audit.utils import log_audit, model_to_dict_safe

from academics.models import ClassLevel, Subject
from students.models import Student

from .forms import ExamForm, ExamResultForm, ExamTypeForm, GradeScaleForm
from .models import Exam, ExamResult, ExamType, GradeScale


# =====================================================
# EXCEL HELPERS
# =====================================================

def style_excel_sheet(ws, title):
    ws.freeze_panes = "A4"

    ws["A1"] = "School Link"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="417690")
    ws["A1"].alignment = Alignment(horizontal="left")

    ws["A2"] = title
    ws["A2"].font = Font(size=12, bold=True, color="2B5668")

    thin = Side(border_style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill = PatternFill("solid", fgColor="417690")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")

    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column

        for cell in column_cells:
            try:
                value_length = len(str(cell.value))
                if value_length > max_length:
                    max_length = value_length
            except Exception:
                pass

        ws.column_dimensions[get_column_letter(column)].width = min(max_length + 4, 40)


def excel_response(workbook, filename):
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def safe_delete_exam_object(request, obj, success_message, redirect_url, audit_message=None):
    old_values = model_to_dict_safe(obj)
    object_repr = str(obj)

    try:
        log_audit(
            request,
            action="delete",
            obj=obj,
            message=audit_message or f"Deleted {object_repr}",
            old_values=old_values,
            new_values={},
        )

        obj.delete()
        messages.success(request, success_message)

    except ProtectedError:
        messages.error(
            request,
            "This record cannot be deleted because it is already used somewhere. "
            "You can edit it or mark it inactive instead."
        )

    except Exception:
        messages.error(
            request,
            "This record could not be deleted. Please check if it is connected to other records."
        )

    return redirect(redirect_url)


def result_audit_message(result, prefix):
    return (
        f"{prefix} result for {result.student.full_name} "
        f"- {result.subject.name} - {result.exam.name}"
    )


# =====================================================
# SETUP UPDATE / DELETE PAGES
# =====================================================

@permission_required("exams.manage")
def exam_type_update(request, pk):
    exam_type = get_object_or_404(ExamType, pk=pk)

    if request.method == "POST":
        old_values = model_to_dict_safe(exam_type)
        form = ExamTypeForm(request.POST, instance=exam_type)

        if form.is_valid():
            exam_type = form.save()

            log_audit(
                request,
                action="update",
                obj=exam_type,
                message=f"Updated exam type {exam_type.name}",
                old_values=old_values,
                new_values=model_to_dict_safe(exam_type),
            )

            messages.success(request, "Exam type updated successfully.")
            return redirect("exam_type_list")

        messages.error(request, "Please correct the exam type form.")
    else:
        form = ExamTypeForm(instance=exam_type)

    return render(request, "exams/simple_exam_form.html", {
        "form": form,
        "title": "Update Exam Type",
        "button_text": "Update Exam Type",
        "back_url": "exam_type_list",
    })


@permission_required("exams.manage")
def exam_type_delete(request, pk):
    exam_type = get_object_or_404(ExamType, pk=pk)

    if request.method == "POST":
        return safe_delete_exam_object(
            request,
            exam_type,
            "Exam type deleted successfully.",
            "exam_type_list",
            audit_message=f"Deleted exam type {exam_type.name}",
        )

    return redirect("exam_type_list")


@permission_required("exams.manage")
def grade_scale_update(request, pk):
    grade = get_object_or_404(GradeScale, pk=pk)

    if request.method == "POST":
        old_values = model_to_dict_safe(grade)
        form = GradeScaleForm(request.POST, instance=grade)

        if form.is_valid():
            grade = form.save()

            log_audit(
                request,
                action="update",
                obj=grade,
                message=f"Updated grade scale {grade.grade}",
                old_values=old_values,
                new_values=model_to_dict_safe(grade),
            )

            messages.success(request, "Grade scale updated successfully.")
            return redirect("grade_scale_list")

        messages.error(request, "Please correct the grade scale form.")
    else:
        form = GradeScaleForm(instance=grade)

    return render(request, "exams/simple_exam_form.html", {
        "form": form,
        "title": "Update Grade Scale",
        "button_text": "Update Grade",
        "back_url": "grade_scale_list",
    })


@permission_required("exams.manage")
def grade_scale_delete(request, pk):
    grade = get_object_or_404(GradeScale, pk=pk)

    if request.method == "POST":
        return safe_delete_exam_object(
            request,
            grade,
            "Grade scale deleted successfully.",
            "grade_scale_list",
            audit_message=f"Deleted grade scale {grade.grade}",
        )

    return redirect("grade_scale_list")


@permission_required("exams.manage")
def exam_update(request, pk):
    exam = get_object_or_404(Exam, pk=pk)

    if request.method == "POST":
        old_values = model_to_dict_safe(exam)
        form = ExamForm(request.POST, instance=exam)

        if form.is_valid():
            updated_exam = form.save()

            log_audit(
                request,
                action="update",
                obj=updated_exam,
                message=f"Updated exam {updated_exam.name}",
                old_values=old_values,
                new_values=model_to_dict_safe(updated_exam),
            )

            messages.success(request, "Exam updated successfully.")
            return redirect("exam_detail", pk=updated_exam.pk)

        messages.error(request, "Please correct the exam form.")
    else:
        form = ExamForm(instance=exam)

    return render(request, "exams/exam_form.html", {
        "form": form,
        "exam": exam,
        "title": "Update Exam",
        "button_text": "Update Exam",
        "is_update": True,
    })


@permission_required("exams.manage")
def exam_delete(request, pk):
    exam = get_object_or_404(Exam, pk=pk)

    if request.method == "POST":
        return safe_delete_exam_object(
            request,
            exam,
            "Exam deleted successfully.",
            "exam_list",
            audit_message=f"Deleted exam {exam.name}",
        )

    return redirect("exam_list")


# =====================================================
# NORMAL EXAM PAGES
# =====================================================

@permission_required("exams.view")
def exams_home(request):
    total_exams = Exam.objects.count()
    open_exams = Exam.objects.filter(status="open").count()
    published_exams = Exam.objects.filter(status="published").count()
    total_results = ExamResult.objects.count()

    recent_exams = Exam.objects.select_related(
        "exam_type",
        "academic_year",
        "term",
        "class_level",
    ).order_by("-created_at")[:8]

    recent_results = ExamResult.objects.select_related(
        "exam",
        "student",
        "subject",
    ).order_by("-created_at")[:8]

    context = {
        "total_exams": total_exams,
        "open_exams": open_exams,
        "published_exams": published_exams,
        "total_results": total_results,
        "recent_exams": recent_exams,
        "recent_results": recent_results,
    }

    return render(request, "exams/home.html", context)


@permission_required("exams.manage")
def exam_type_list(request):
    exam_types = ExamType.objects.all()

    if request.method == "POST":
        form = ExamTypeForm(request.POST)

        if form.is_valid():
            exam_type = form.save()

            log_audit(
                request,
                action="create",
                obj=exam_type,
                message=f"Created exam type {exam_type.name}",
                old_values={},
                new_values=model_to_dict_safe(exam_type),
            )

            messages.success(request, "Exam type saved successfully.")
            return redirect("exam_type_list")

        messages.error(request, "Please correct the exam type form.")
    else:
        form = ExamTypeForm()

    return render(request, "exams/exam_type_list.html", {
        "exam_types": exam_types,
        "form": form,
    })


@permission_required("exams.manage")
def grade_scale_list(request):
    grades = GradeScale.objects.all()

    if request.method == "POST":
        form = GradeScaleForm(request.POST)

        if form.is_valid():
            grade = form.save()

            log_audit(
                request,
                action="create",
                obj=grade,
                message=f"Created grade scale {grade.grade}",
                old_values={},
                new_values=model_to_dict_safe(grade),
            )

            messages.success(request, "Grade scale saved successfully.")
            return redirect("grade_scale_list")

        messages.error(request, "Please correct the grade scale form.")
    else:
        form = GradeScaleForm()

    return render(request, "exams/grade_scale_list.html", {
        "grades": grades,
        "form": form,
    })


@permission_required("exams.view")
def exam_list(request):
    query = request.GET.get("q", "").strip()
    class_filter = request.GET.get("class", "").strip()
    status_filter = request.GET.get("status", "").strip()

    exams = Exam.objects.select_related(
        "exam_type",
        "academic_year",
        "term",
        "class_level",
    ).all()

    if query:
        exams = exams.filter(
            Q(name__icontains=query) |
            Q(exam_type__name__icontains=query)
        )

    if class_filter:
        exams = exams.filter(class_level_id=class_filter)

    if status_filter:
        exams = exams.filter(status=status_filter)

    context = {
        "exams": exams,
        "query": query,
        "class_filter": class_filter,
        "status_filter": status_filter,
        "classes": ClassLevel.objects.filter(is_active=True),
        "status_choices": Exam.STATUS_CHOICES,
    }

    return render(request, "exams/exam_list.html", context)


@permission_required("exams.manage")
def exam_create(request):
    if request.method == "POST":
        form = ExamForm(request.POST)

        if form.is_valid():
            exam = form.save()

            log_audit(
                request,
                action="create",
                obj=exam,
                message=f"Created exam {exam.name}",
                old_values={},
                new_values=model_to_dict_safe(exam),
            )

            messages.success(request, "Exam created successfully.")
            return redirect("exam_detail", pk=exam.pk)

        messages.error(request, "Please correct the exam form.")
    else:
        form = ExamForm()

    return render(request, "exams/exam_form.html", {
        "form": form,
        "title": "Create Exam",
    })


@permission_required("exams.view")
def exam_detail(request, pk):
    exam = get_object_or_404(
        Exam.objects.select_related(
            "exam_type",
            "academic_year",
            "term",
            "class_level",
        ),
        pk=pk
    )

    results = ExamResult.objects.select_related(
        "student",
        "subject",
        "teacher",
    ).filter(exam=exam)

    total_students = Student.objects.filter(
        class_level=exam.class_level,
        status="active"
    ).count()

    students_with_results = results.values("student").distinct().count()
    subject_count = results.values("subject").distinct().count()

    average_marks = results.aggregate(avg=Avg("marks"))["avg"] or Decimal("0.00")

    context = {
        "exam": exam,
        "results": results,
        "total_students": total_students,
        "students_with_results": students_with_results,
        "subject_count": subject_count,
        "average_marks": average_marks,
    }

    return render(request, "exams/exam_detail.html", context)


# =====================================================
# RESULTS
# =====================================================

@permission_required("results.view")
def result_list(request):
    query = request.GET.get("q", "").strip()
    exam_filter = request.GET.get("exam", "").strip()
    class_filter = request.GET.get("class", "").strip()
    subject_filter = request.GET.get("subject", "").strip()

    results = ExamResult.objects.select_related(
        "exam",
        "exam__class_level",
        "student",
        "student__class_level",
        "student__stream",
        "subject",
        "teacher",
    ).all()

    if query:
        results = results.filter(
            Q(student__admission_number__icontains=query) |
            Q(student__first_name__icontains=query) |
            Q(student__middle_name__icontains=query) |
            Q(student__last_name__icontains=query)
        )

    if exam_filter:
        results = results.filter(exam_id=exam_filter)

    if class_filter:
        results = results.filter(student__class_level_id=class_filter)

    if subject_filter:
        results = results.filter(subject_id=subject_filter)

    grouped = {}

    for result in results:
        key = (result.exam_id, result.student_id)

        if key not in grouped:
            grouped[key] = {
                "exam": result.exam,
                "student": result.student,
                "subject_count": 0,
                "total_marks": Decimal("0.00"),
                "average_marks": Decimal("0.00"),
                "position": None,
                "total_students": 0,
            }

        grouped[key]["subject_count"] += 1
        grouped[key]["total_marks"] += result.marks

    result_rows = list(grouped.values())

    for row in result_rows:
        if row["subject_count"] > 0:
            row["average_marks"] = row["total_marks"] / row["subject_count"]

    exams_in_rows = {}

    for row in result_rows:
        exam_id = row["exam"].id
        exams_in_rows.setdefault(exam_id, []).append(row)

    for exam_id, rows in exams_in_rows.items():
        rows.sort(key=lambda item: item["average_marks"], reverse=True)

        total_students = Student.objects.filter(
            class_level=rows[0]["exam"].class_level,
            status="active"
        ).count()

        for index, row in enumerate(rows, start=1):
            row["position"] = index
            row["total_students"] = total_students

    result_rows.sort(
        key=lambda item: (
            item["exam"].name,
            item["position"] or 999999,
            item["student"].first_name,
        )
    )

    context = {
        "result_rows": result_rows,
        "query": query,
        "exam_filter": exam_filter,
        "class_filter": class_filter,
        "subject_filter": subject_filter,
        "exams": Exam.objects.filter(is_active=True),
        "classes": ClassLevel.objects.filter(is_active=True),
        "subjects": Subject.objects.filter(is_active=True),
    }

    return render(request, "exams/result_list.html", context)


@permission_required("results.manage")
def result_create(request):
    if request.method == "POST":
        form = ExamResultForm(request.POST)

        if form.is_valid():
            result = form.save()

            log_audit(
                request,
                action="create",
                obj=result,
                message=result_audit_message(result, "Created"),
                old_values={},
                new_values=model_to_dict_safe(result),
            )

            messages.success(request, "Exam result saved successfully.")
            return redirect("student_report_card", student_id=result.student.pk)

        messages.error(request, "Please correct the result form.")
    else:
        form = ExamResultForm()

    return render(request, "exams/result_form.html", {
        "form": form,
        "title": "Enter Result",
        "button_text": "Save Result",
        "is_update": False,
    })


@permission_required("results.manage")
def result_update(request, pk):
    result = get_object_or_404(
        ExamResult.objects.select_related("student", "exam", "subject"),
        pk=pk
    )

    if request.method == "POST":
        old_values = model_to_dict_safe(result)
        form = ExamResultForm(request.POST, instance=result)

        if form.is_valid():
            updated_result = form.save()

            log_audit(
                request,
                action="update",
                obj=updated_result,
                message=result_audit_message(updated_result, "Updated"),
                old_values=old_values,
                new_values=model_to_dict_safe(updated_result),
            )

            messages.success(request, "Exam result updated successfully.")
            return redirect(
                "student_report_card",
                student_id=updated_result.student.pk
            )

        messages.error(request, "Please correct the result form.")
    else:
        form = ExamResultForm(instance=result)

    return render(request, "exams/result_form.html", {
        "form": form,
        "result": result,
        "title": "Update Result",
        "button_text": "Update Result",
        "is_update": True,
    })


@permission_required("results.manage")
def result_delete(request, pk):
    result = get_object_or_404(
        ExamResult.objects.select_related("student", "exam", "subject"),
        pk=pk
    )

    student_id = result.student_id
    old_values = model_to_dict_safe(result)

    if request.method == "POST":
        log_audit(
            request,
            action="delete",
            obj=result,
            message=result_audit_message(result, "Deleted"),
            old_values=old_values,
            new_values={},
        )

        result.delete()
        messages.success(request, "Exam result deleted successfully.")
        return redirect("student_report_card", student_id=student_id)

    return redirect("student_report_card", student_id=student_id)


@permission_required("results.view")
def student_report_card(request, student_id):
    student = get_object_or_404(
        Student.objects.select_related(
            "class_level",
            "stream",
            "parent_guardian",
        ),
        pk=student_id
    )

    exam_id = request.GET.get("exam", "").strip()

    results = ExamResult.objects.select_related(
        "exam",
        "subject",
        "teacher",
    ).filter(student=student)

    selected_exam = None
    position = None
    total_students = 0
    subject_count = 0

    if exam_id:
        selected_exam = get_object_or_404(Exam, pk=exam_id)
        results = results.filter(exam=selected_exam)

        total_students = Student.objects.filter(
            class_level=selected_exam.class_level,
            status="active"
        ).count()

        exam_results = ExamResult.objects.select_related("student").filter(
            exam=selected_exam,
            student__class_level=selected_exam.class_level,
            student__status="active",
        )

        grouped_results = defaultdict(list)

        for result in exam_results:
            grouped_results[result.student_id].append(result)

        performance_rows = []

        for grouped_student_id, student_results in grouped_results.items():
            total = sum([item.marks for item in student_results], Decimal("0.00"))
            count = len(student_results)
            average = total / count if count else Decimal("0.00")

            performance_rows.append({
                "student_id": grouped_student_id,
                "total": total,
                "average": average,
                "count": count,
            })

        performance_rows.sort(key=lambda row: row["average"], reverse=True)

        for index, row in enumerate(performance_rows, start=1):
            if row["student_id"] == student.id:
                position = index
                subject_count = row["count"]
                break

    total_marks = results.aggregate(total=Sum("marks"))["total"] or Decimal("0.00")
    average_marks = results.aggregate(avg=Avg("marks"))["avg"] or Decimal("0.00")

    context = {
        "student": student,
        "results": results,
        "exams": Exam.objects.filter(is_active=True),
        "selected_exam": selected_exam,
        "exam_id": exam_id,
        "total_marks": total_marks,
        "average_marks": average_marks,
        "position": position,
        "total_students": total_students,
        "subject_count": subject_count,
    }

    return render(request, "exams/student_report_card.html", context)


@permission_required("results.manage")
def bulk_marks_entry(request):
    selected_exam_id = request.GET.get("exam") or request.POST.get("exam") or ""
    selected_subject_id = request.GET.get("subject") or request.POST.get("subject") or ""

    exams = Exam.objects.select_related("class_level").filter(is_active=True)
    subjects = Subject.objects.filter(is_active=True)

    selected_exam = None
    selected_subject = None
    students = Student.objects.none()
    student_rows = []

    if selected_exam_id:
        selected_exam = get_object_or_404(Exam, pk=selected_exam_id)

        students = Student.objects.select_related(
            "class_level",
            "stream"
        ).filter(
            class_level=selected_exam.class_level,
            status="active"
        ).order_by("stream__name", "first_name", "last_name")

    if selected_subject_id:
        selected_subject = get_object_or_404(Subject, pk=selected_subject_id)

    existing_results = {}

    if selected_exam and selected_subject:
        existing_results = {
            result.student_id: result
            for result in ExamResult.objects.filter(
                exam=selected_exam,
                subject=selected_subject
            )
        }

    if request.method == "POST":
        if not selected_exam or not selected_subject:
            messages.error(request, "Please choose exam and subject first.")
            return redirect("bulk_marks_entry")

        saved_count = 0
        created_count = 0
        updated_count = 0

        for student in students:
            marks_value = request.POST.get(f"marks_{student.id}", "").strip()

            if marks_value == "":
                continue

            try:
                marks = Decimal(marks_value)
            except Exception:
                marks = Decimal("0.00")

            if marks < 0:
                marks = Decimal("0.00")

            if marks > 100:
                marks = Decimal("100.00")

            old_result = existing_results.get(student.id)
            old_values = model_to_dict_safe(old_result) if old_result else {}

            result, created = ExamResult.objects.update_or_create(
                exam=selected_exam,
                student=student,
                subject=selected_subject,
                defaults={
                    "marks": marks,
                }
            )

            if created:
                created_count += 1
                audit_action = "create"
                audit_prefix = "Created by bulk entry"
            else:
                updated_count += 1
                audit_action = "update"
                audit_prefix = "Updated by bulk entry"

            log_audit(
                request,
                action=audit_action,
                obj=result,
                message=result_audit_message(result, audit_prefix),
                old_values=old_values,
                new_values=model_to_dict_safe(result),
            )

            saved_count += 1

        log_audit(
            request,
            action="generate",
            obj=selected_exam,
            message=(
                f"Bulk marks saved for {saved_count} students. "
                f"Exam: {selected_exam.name}, Subject: {selected_subject.name}, "
                f"Created: {created_count}, Updated: {updated_count}"
            ),
            old_values={},
            new_values={
                "exam": selected_exam.name,
                "subject": selected_subject.name,
                "saved_count": saved_count,
                "created_count": created_count,
                "updated_count": updated_count,
            },
        )

        messages.success(request, f"Marks saved for {saved_count} students.")
        return redirect("result_list")

    for student in students:
        result = existing_results.get(student.id)

        student_rows.append({
            "student": student,
            "result": result,
            "marks": result.marks if result else "",
            "grade": result.grade if result else "",
            "remark": result.remark if result else "",
        })

    context = {
        "exams": exams,
        "subjects": subjects,
        "selected_exam_id": selected_exam_id,
        "selected_subject_id": selected_subject_id,
        "selected_exam": selected_exam,
        "selected_subject": selected_subject,
        "student_rows": student_rows,
    }

    return render(request, "exams/bulk_marks_entry.html", context)


# =====================================================
# REPORTS
# =====================================================

@permission_required("results.reports")
def class_performance_report(request):
    exam_id = request.GET.get("exam", "").strip()

    exams = Exam.objects.select_related(
        "class_level",
        "exam_type",
        "academic_year",
        "term",
    ).filter(is_active=True)

    selected_exam = None
    performance_rows = []

    if exam_id:
        selected_exam = get_object_or_404(Exam, pk=exam_id)

        students = Student.objects.select_related(
            "class_level",
            "stream"
        ).filter(
            class_level=selected_exam.class_level,
            status="active"
        ).order_by("stream__name", "first_name", "last_name")

        results = ExamResult.objects.select_related(
            "student",
            "subject"
        ).filter(exam=selected_exam)

        grouped_results = defaultdict(list)

        for result in results:
            grouped_results[result.student_id].append(result)

        for student in students:
            student_results = grouped_results.get(student.id, [])

            total_marks = sum([r.marks for r in student_results], Decimal("0.00"))
            subject_count = len(student_results)
            average_marks = total_marks / subject_count if subject_count else Decimal("0.00")

            performance_rows.append({
                "student": student,
                "subject_count": subject_count,
                "total_marks": total_marks,
                "average_marks": average_marks,
            })

        performance_rows.sort(key=lambda row: row["average_marks"], reverse=True)

        for index, row in enumerate(performance_rows, start=1):
            row["position"] = index

    context = {
        "exams": exams,
        "exam_id": exam_id,
        "selected_exam": selected_exam,
        "performance_rows": performance_rows,
    }

    return render(request, "exams/class_performance_report.html", context)


# =====================================================
# EXAM EXCEL EXPORTS
# =====================================================

@permission_required("reports.exams")
def export_results_excel(request):
    query = request.GET.get("q", "").strip()
    exam_filter = request.GET.get("exam", "").strip()
    class_filter = request.GET.get("class", "").strip()
    subject_filter = request.GET.get("subject", "").strip()

    results = ExamResult.objects.select_related(
        "exam",
        "student",
        "student__class_level",
        "student__stream",
        "subject",
        "teacher",
    ).all()

    if query:
        results = results.filter(
            Q(student__admission_number__icontains=query) |
            Q(student__first_name__icontains=query) |
            Q(student__middle_name__icontains=query) |
            Q(student__last_name__icontains=query)
        )

    if exam_filter:
        results = results.filter(exam_id=exam_filter)

    if class_filter:
        results = results.filter(student__class_level_id=class_filter)

    if subject_filter:
        results = results.filter(subject_id=subject_filter)

    log_audit(
        request,
        action="export",
        obj=None,
        app_label="exams",
        model_name="examresult",
        object_repr="Exam Results Excel",
        message="Exported exam results Excel report.",
        old_values={},
        new_values={
            "query": query,
            "exam_filter": exam_filter,
            "class_filter": class_filter,
            "subject_filter": subject_filter,
            "records_count": results.count(),
        },
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Exam Results"

    headers = [
        "#",
        "Admission No",
        "Student",
        "Class",
        "Stream",
        "Exam",
        "Subject",
        "Marks",
        "Grade",
        "Points",
        "Remark",
        "Teacher",
    ]

    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(headers)

    for index, result in enumerate(results, start=1):
        ws.append([
            index,
            result.student.admission_number,
            result.student.full_name,
            result.student.class_level.name,
            result.student.stream.name if result.student.stream else "",
            result.exam.name,
            result.subject.name,
            float(result.marks),
            result.grade,
            float(result.points),
            result.remark,
            result.teacher.full_name if result.teacher else "",
        ])

    style_excel_sheet(ws, "Exam Results Report")

    return excel_response(wb, "exam_results.xlsx")


@permission_required("reports.exams")
def export_class_performance_excel(request):
    exam_id = request.GET.get("exam", "").strip()

    wb = Workbook()
    ws = wb.active
    ws.title = "Class Performance"

    headers = [
        "Position",
        "Admission No",
        "Student",
        "Class",
        "Stream",
        "Subjects",
        "Total Marks",
        "Average",
    ]

    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(headers)

    exported_rows = 0
    selected_exam_name = ""

    if exam_id:
        selected_exam = get_object_or_404(Exam, pk=exam_id)
        selected_exam_name = selected_exam.name

        students = Student.objects.select_related(
            "class_level",
            "stream"
        ).filter(
            class_level=selected_exam.class_level,
            status="active"
        ).order_by("stream__name", "first_name", "last_name")

        results = ExamResult.objects.select_related(
            "student",
            "subject"
        ).filter(exam=selected_exam)

        grouped_results = defaultdict(list)

        for result in results:
            grouped_results[result.student_id].append(result)

        performance_rows = []

        for student in students:
            student_results = grouped_results.get(student.id, [])

            total_marks = sum([r.marks for r in student_results], Decimal("0.00"))
            subject_count = len(student_results)
            average_marks = total_marks / subject_count if subject_count else Decimal("0.00")

            performance_rows.append({
                "student": student,
                "subject_count": subject_count,
                "total_marks": total_marks,
                "average_marks": average_marks,
            })

        performance_rows.sort(key=lambda row: row["average_marks"], reverse=True)

        for index, row in enumerate(performance_rows, start=1):
            student = row["student"]

            ws.append([
                index,
                student.admission_number,
                student.full_name,
                student.class_level.name,
                student.stream.name if student.stream else "",
                row["subject_count"],
                float(row["total_marks"]),
                float(row["average_marks"]),
            ])

            exported_rows += 1

    log_audit(
        request,
        action="export",
        obj=None,
        app_label="exams",
        model_name="examresult",
        object_repr="Class Performance Excel",
        message="Exported class performance Excel report.",
        old_values={},
        new_values={
            "exam_id": exam_id,
            "exam": selected_exam_name,
            "records_count": exported_rows,
        },
    )

    style_excel_sheet(ws, "Class Performance Report")

    return excel_response(wb, "class_performance.xlsx")